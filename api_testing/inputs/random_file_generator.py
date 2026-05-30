import io
import random
from typing import Any, Literal, Callable
from faker import Faker

from api_testing.inputs.fuzz_strategy import FuzzStrategy
from .random_generator import RandomGenerator


def _pil_image_generator(width=800, height=200, image_format="PNG"):
    try:
        from PIL import Image, ImageDraw, ImageFont
        img = Image.new("RGB", (width, height), color=(73, 109, 137))
        d = ImageDraw.Draw(img)
        text = "Sample Generated Image"
        try:
            font = ImageFont.truetype("arial.ttf", 20)
        except Exception:
            font = ImageFont.load_default()
        d.text((width // 2 - 60, height // 2 - 10), text, fill=(255, 255, 255), font=font)
        buf = io.BytesIO()
        img.save(buf, format=image_format)
        return buf.getvalue()
    except Exception:
        return b""


class RandomFileGenerator(RandomGenerator):
    """
    Generate random file data for upload testing.

    Return format:
        (filename: str, content: bytes, content_type: str)
    """

    SUPPORTED_TYPES = ("pdf", "docx", "txt", "png", "jpeg", "bmp")

    MIME_TYPES = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "txt": "text/plain",
        "png": "image/png",
        "jpeg": "image/jpeg",
        "bmp": "image/bmp",
    }

    EXTENSIONS = {
        "pdf": ".pdf",
        "docx": ".docx",
        "txt": ".txt",
        "png": ".png",
        "jpeg": ".jpeg",
        "bmp": ".bmp",
    }
    description: str = """
    A generator that produces random file data for upload testing.
    Attributes:
        file_type (Literal["pdf", "docx", "txt", "png", "jpeg", "bmp"] | None):
            Optional fixed file type to generate. If None, a random supported type is selected.
    """
    def __init__(
        self,
        file_type: Literal["pdf", "docx", "txt", "png", "jpeg", "bmp"] = None,
        seed: int | None = None,
    ):
        super().__init__(seed)

        self.file_type = file_type.lower() if file_type else None
        if self.file_type and self.file_type not in self.SUPPORTED_TYPES:
            raise ValueError(f"Unsupported file type: {file_type}")

        self.fake = Faker()
        self._register_providers()

        self._generators: dict[str, Callable[[], tuple[str, bytes, str]]] = {
            t: lambda t=t: self._generate_file(t)
            for t in self.SUPPORTED_TYPES
        }

    def _register_providers(self):
        pass

    def _generate_file(self, file_type: str) -> tuple[str, bytes, str]:
        method_name = f"{file_type}_file"

        try:
            content = getattr(self.fake, method_name)(raw=True)
        except Exception:
            content = self._fallback_generate(file_type)

        filename = f"{self.fake.uuid4()}{self.EXTENSIONS[file_type]}"
        content_type = self.MIME_TYPES[file_type]
        return (filename, content, content_type)

    def _fallback_generate(self, file_type: str) -> bytes:
        if file_type == "txt":
            return self.fake.text(max_nb_chars=500).encode("utf-8")
        elif file_type == "png":
            return _pil_image_generator(image_format="PNG")
        elif file_type == "jpeg":
            return _pil_image_generator(image_format="JPEG")
        elif file_type == "bmp":
            return _pil_image_generator(image_format="BMP")
        elif file_type == "pdf":
            return self._fallback_pdf()
        elif file_type == "docx":
            return self._fallback_docx()
        return self.rand.randbytes(1024)

    def _fallback_pdf(self) -> bytes:
        try:
            from faker import Faker as FakeFaker
            f = FakeFaker()
            try:
                from reportlab.pdfgen import canvas
                from reportlab.lib.pagesizes import letter
                pdf_buf = io.BytesIO()
                c = canvas.Canvas(pdf_buf, pagesize=letter)
                text_obj = c.beginText(50, 742)
                text_obj.setFont("Helvetica", 12)
                for line in f.text(max_nb_chars=500).split("\n"):
                    text_obj.textLine(line)
                c.drawText(text_obj)
                c.save()
                return pdf_buf.getvalue()
            except Exception:
                pass
            try:
                from faker_file.providers.pdf_file.generators.pil_generator import PilPdfGenerator
                gen = PilPdfGenerator(generator=f)
                content = gen.generate("Sample PDF content", {}, f)
                return content
            except Exception:
                pass
            return self.rand.randbytes(512)
        except Exception:
            return b"PDF content"

    def _fallback_docx(self) -> bytes:
        try:
            buf = io.BytesIO()
            from docx import Document
            doc = Document()
            doc.add_heading("Sample Document", 0)
            doc.add_paragraph(self.fake.text(max_nb_chars=500))
            doc.save(buf)
            return buf.getvalue()
        except Exception:
            return b"Docx content"

    # --------------------------
    # Public API
    # --------------------------
    def next_value(self, context_pool=None, *args, **kwargs) -> tuple[str, bytes, str]:
        file_type = self.file_type
        if  file_type is None:
            file_type = random.choice(self.SUPPORTED_TYPES)
        return self._generators[file_type]()

    # --------------------------
    # Fuzzing
    # --------------------------
    def next_fuzz_value(
        self,
        strategy: FuzzStrategy | None,
        context_pool=None,
        *args,
        **kwargs,
    ) -> Any:

        strategies = [
            FuzzStrategy.EMPTY,
            FuzzStrategy.LARGE,
            FuzzStrategy.WRONG_TYPE,
            FuzzStrategy.JUNK,
            FuzzStrategy.CORRUPT,
        ]

        if strategy is None:
            strategy = self.rand.choice(strategies)

        filename, content, content_type = self.next_value()

        if strategy == FuzzStrategy.EMPTY:
            return (filename, b"", content_type)

        if strategy == FuzzStrategy.LARGE:
            return (filename, b"A" * (10 * 1024 * 1024), content_type)

        if strategy == FuzzStrategy.WRONG_TYPE:
            return f"invalid_file_{self.rand.getrandbits(32)}"

        if strategy == FuzzStrategy.JUNK:
            return (filename, self.rand.randbytes(2048), content_type)

        if strategy == FuzzStrategy.CORRUPT:
            if not content:
                return (filename, b"\xFF\xFF", content_type)

            mutable = bytearray(content)
            for i in range(min(16, len(mutable))):
                mutable[i] = self.rand.randint(0, 255)

            return (filename, bytes(mutable), content_type)

        return None